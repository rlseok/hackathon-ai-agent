"""라인 센싱 시뮬레이션 결과(results/live_sim/summary.json) → 정확도 표.

  results/live_sim/line_sim_accuracy.xlsx   Summary · By defect type · Definitions (영어, 팀 공유용)
  results/live_sim/README.md                같은 표 (한국어)

실행: .venv\\Scripts\\python tools/live_sim_table.py   (먼저 tools/record_line_sim.py)
"""
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from agent.i18n import type_name  # noqa: E402
from agent.knowledge import lookup  # noqa: E402

DIR = ROOT / 'results/live_sim'
ORDER = ['pcb', 'phone', 'metal']
LINE_KO = {'pcb': 'PCB', 'phone': '액정(휴대폰 표면)', 'metal': '금속 표면'}
HEAD = PatternFill('solid', fgColor='1F4E79')
ZEBRA = PatternFill('solid', fgColor='EEF3F8')
THIN = Side(style='thin', color='B4C6E7')


def pct(x):
    return f'{100 * x:.1f}%' if x is not None else '-'


def tname(t, domain, lang):
    if t == 'normal':
        return {'ko': '정상', 'en': 'Normal'}[lang]
    types = [x.replace('_', ' ') for x in t.replace('Dark Spot', 'Dark_Spot').split()]   # 대조 기록은 유형을 공백으로 이어 둔다
    kb = {k['type']: k for k in lookup(types, domain)}
    return ' + '.join(type_name(x, kb.get(x), lang) for x in types)


def by_type(x):
    return sorted(x['by_type'].items(), key=lambda kv: (kv[0] == 'normal', kv[0]))   # 정상은 맨 뒤


def sheet(ws, header, rows, widths):
    ws.append(header)
    for r in rows:
        ws.append(r)
    for c, w in zip('ABCDEFGHIJKLMNOP', widths):
        ws.column_dimensions[c].width = w
    for i, row in enumerate(ws.iter_rows(), 1):
        for cell in row:
            cell.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
            cell.alignment = Alignment(vertical='center', wrap_text=True)
            if i == 1:
                cell.fill, cell.font = HEAD, Font(bold=True, color='FFFFFF')
            elif i % 2 == 0:
                cell.fill = ZEBRA
    ws.freeze_panes = 'A2'


def main():
    s = json.loads((DIR / 'summary.json').read_text(encoding='utf-8'))
    lines = [s[k] for k in ORDER if k in s]
    wb = Workbook()
    ws = wb.active
    ws.title = 'Summary'
    sheet(ws, ['Line', 'Video length', 'Items in video', 'Inspected in sensing zone', 'Product verdict accuracy',
               'Defective items flagged', 'Normal items passed', 'Abnormal / Suspect / Normal verdicts',
               'Defect TP / FP / FN', 'Precision', 'Recall', 'F1', 'Processing ms/frame (mean / p95)', 'Processing capacity'],
          [[x['name'], f'{x["video_seconds"]} s', x['items_in_video'], x['inspected'], pct(x['board_acc']),
            f'{x["defect_caught"]}/{x["defect_boards"]}', f'{x["normal_passed"]}/{x["normal_boards"]}',
            f'{x["verdicts"]["reject"]} / {x["verdicts"]["review"]} / {x["verdicts"]["normal"]}',
            f'{x["tp"]} / {x["fp"]} / {x["fn"]}', round(x['precision'], 3), round(x['recall'], 3), round(x['f1'], 3),
            f'{x["proc_ms_mean"]} / {x["proc_ms_p95"]}', f'{x["proc_fps"]} fps'] for x in lines],
          [18, 12, 12, 14, 14, 14, 13, 20, 16, 10, 10, 8, 18, 13])
    ws2 = wb.create_sheet('By defect type')
    rows = []
    for x in lines:
        for t, v in by_type(x):
            rows.append([x['name'], tname(t, x['domain'], 'en'), v['items'], v['correct'], pct(v['correct'] / v['items']),
                         v['verdicts'].get('reject', 0), v['verdicts'].get('review', 0), v['verdicts'].get('normal', 0)])
    sheet(ws2, ['Line', 'Ground-truth type', 'Items', 'Verdict correct', 'Accuracy', 'Abnormal', 'Suspect', 'Normal'],
          rows, [18, 30, 8, 14, 10, 10, 10, 10])
    ws3 = wb.create_sheet('Definitions')
    sheet(ws3, ['Term', 'Meaning'], [
        ['Simulation', 'Each virtual conveyor video is processed frame by frame with the same code as the live video server '
                       '(no frames skipped), so every item passes the sensing zone once. No AI report text is written.'],
        ['Product verdict accuracy', 'A defective item counts as correct if it is flagged Abnormal or Suspect; a normal item '
                                     'counts as correct if it is judged Normal.'],
        ['Defect TP / FP / FN', 'Per detected defect: same type and IoU ≥ 0.5 with a ground-truth label = TP; extra detections '
                                '= FP; missed labels = FN.'],
        ['Precision / Recall / F1', 'TP/(TP+FP), TP/(TP+FN), and their harmonic mean.'],
        ['Suspect by rule', 'Some types are at most Suspect by the verdict rules (PCB scratch and pinhole, phone scratch and '
                            'stain, local metal porosity and scratch, …), so Suspect is the expected verdict for them.'],
        ['Processing ms/frame', 'Time for one frame: product finding, sensing-zone inspection when an item passes, '
                                'full-frame preview detection and drawing (RTX 5060 Laptop GPU). Video encoding excluded.'],
    ], [26, 110])
    wb.save(DIR / 'line_sim_accuracy.xlsx')

    md = ['# 라인 센싱 시뮬레이션 결과', '',
          '가상 컨베이어 영상을 영상 서버와 같은 처리로 한 프레임도 건너뛰지 않고 돌려, 모든 제품이 센싱 구역을 한 번씩 지나게 한 결과다.',
          '영상: `videos/<line>_sensing_sim.mp4` (git 제외) · 제품별 대조: `<line>_boards.csv` · 만든 도구: `tools/record_line_sim.py`, `tools/live_sim_table.py`', '',
          '| 라인 | 영상 길이 | 제품 수 | 판정한 제품 | 제품 판정 정확도 | 결함 제품 잡음 | 정상 통과 | 이상/의심/정상 | 결함 단위 P / R / F1 | 처리 (ms/프레임) |',
          '|---|---|---|---|---|---|---|---|---|---|']
    for x in lines:
        md.append(f'| {LINE_KO[x["line"]]} | {x["video_seconds"]}초 | {x["items_in_video"]} | {x["inspected"]} | {pct(x["board_acc"])} | '
                  f'{x["defect_caught"]}/{x["defect_boards"]} | {x["normal_passed"]}/{x["normal_boards"]} | '
                  f'{x["verdicts"]["reject"]}/{x["verdicts"]["review"]}/{x["verdicts"]["normal"]} | '
                  f'{x["precision"]:.3f} / {x["recall"]:.3f} / {x["f1"]:.3f} | {x["proc_ms_mean"]} (p95 {x["proc_ms_p95"]}) |')
    md += ['', '## 정답 유형별', '', '| 라인 | 정답 유형 | 제품 수 | 판정 맞음 | 정확도 | 이상 | 의심 | 정상 |', '|---|---|---|---|---|---|---|---|']
    for x in lines:
        for t, v in by_type(x):
            md.append(f'| {LINE_KO[x["line"]]} | {tname(t, x["domain"], "ko")} | {v["items"]} | {v["correct"]} | '
                      f'{pct(v["correct"] / v["items"])} | {v["verdicts"].get("reject", 0)} | {v["verdicts"].get("review", 0)} | '
                      f'{v["verdicts"].get("normal", 0)} |')
    md += ['', '- 제품 판정 정확도: 결함 제품은 이상·의심으로, 정상 제품은 정상으로 판정하면 맞음.',
           '- 결함 단위: 같은 유형 + IoU 0.5 이상이면 TP. 일부 유형은 규칙상 최대 의심이라 의심이 정상 결과다.',
           '- 처리 시간: 제품 찾기 + 센싱 구역 검사 + 화면 전체 미리보기 검출 + 그리기 (RTX 5060 노트북, 녹화 시간 제외).']
    (DIR / 'README.md').write_text('\n'.join(md) + '\n', encoding='utf-8')
    print('\n'.join(md))


if __name__ == '__main__':
    main()
