"""검사 기록: 검사 대장(엑셀) + 검사 1건당 md 요약.

- outputs/records.jsonl  원본 기록. 한 줄 = 기록 1회, 같은 id 는 마지막 줄이 최신 (작업자 확인으로 상태가 바뀌면 줄 추가)
- outputs/inspection_ledger.xlsx  위 기록으로 매번 다시 만든다. 시트: Inspections(썸네일·판정 색) / Defects / Summary
- outputs/records/<날짜>/<id>.md  사람이 읽는 요약. 대화창에서 과거 사례를 찾을 때도 쓴다
엑셀 대장은 어디서나 읽을 수 있게 영어로 고정하고, md 는 검사 결과의 언어로 쓴다.
"""
import json
import os
import time
from collections import Counter
from pathlib import Path

from PIL import Image

from agent.detect import TAXONOMIES
from agent.i18n import (DOMAIN, PROCESS_NAME, R, STATUS, pos_text, process_name, reason_text, situation_text,
                        type_name, verdict_name)

JSONL = 'records.jsonl'
XLSX = 'inspection_ledger.xlsx'
MD_DIR = 'records'
FILL = {'normal': 'C8E6C9', 'review': 'FFE0B2', 'reject': 'FFCDD2'}
THUMB_H = 60


def _rel(path, start):
    return Path(os.path.relpath(path, start)).as_posix() if path else ''


def make_entry(result, status, out_root):
    det, sit = result['detection'], result['situation']
    kb_map = {k['type']: k for k in result['knowledge']}
    meta = result.get('meta') or {}
    rep = sit.get('representative')
    defects = det['defects']
    n_rej = sum(d['level'] == 'reject' for d in defects)
    return {
        'id': result['name'],
        # 언어와 무관한 값 (엑셀 대장은 이 값으로 영어를 만든다)
        'domain': det['domain'], 'rep': rep or '', 'process': meta.get('process', ''),
        '일시': time.strftime('%Y-%m-%d %H:%M:%S'),
        '도메인': {'pcb': 'PCB', 'msdd': '금속'}.get(det['domain'], det['domain']),
        '대상': det.get('part') or Path(det['image']).name,
        '사진 수': len(det.get('views') or [1]),
        '품번': meta.get('part_no', ''),
        '공정': process_name(meta['process'], 'ko') if meta.get('process') else '',
        '로트': meta.get('lot', ''),
        '판정': verdict_name(det['verdict'], 'ko'),
        'verdict': det['verdict'],
        '대표 유형': type_name(rep, kb_map.get(rep), 'ko') if rep else '',
        '이상': n_rej,
        '의심': len(defects) - n_rej,
        '작업자 수정': sum(d.get('reason_key') == 'edited' for d in defects) + len(det.get('excluded', [])),
        '처리 상태': STATUS['ko'][status],
        'status': status,
        '보고서': _rel(result.get('report'), out_root),
        'md': '',
        'image': det.get('marked_image') or det['image'],
        'result_json': _rel(Path(result['work']) / f'{result["name"]}_result.json', out_root),
        'defects': [{'no': d['id'], '사진': d.get('view', ''), '유형': type_name(d['type'], kb_map.get(d['type']), 'ko'),
                     '클래스': d['type'], '판정': verdict_name(d['level'], 'ko'), '신뢰도': d['conf'],
                     '위치': d['position'], '크기(px)': f'{d["size_px"][0]}×{d["size_px"][1]}',
                     '면적(%)': round(100 * d['area_ratio'], 2) if 'area_ratio' in d else '',
                     '판정 근거': reason_text(d, 'ko'),
                     'level': d['level'], 'reason_key': d.get('reason_key'), 'reason_args': d.get('reason_args', {})}
                    for d in defects],
    }


def write_md(result, entry, out_root):
    """검사 1건 요약 md (검사 결과의 언어)."""
    lang = result.get('lang', 'ko')
    det, kb_map = result['detection'], {k['type']: k for k in result['knowledge']}
    sit = situation_text(result['situation'], det['verdict'], lang)
    path = Path(out_root) / MD_DIR / entry['일시'][:10].replace('-', '') / f'{entry["id"]}.md'
    path.parent.mkdir(parents=True, exist_ok=True)
    meta = result.get('meta') or {}
    tx = result['text']
    lines = [f'# {entry["id"]}', '',
             f'- {R("date", lang)}: {entry["일시"]}',
             f'- {R("target", lang)}: {DOMAIN[lang].get(det["domain"], det["domain"])} / {entry["대상"]} ({entry["사진 수"]})',
             f'- {R("part_info", lang)}: {meta.get("part_no") or "-"} / {process_name(meta.get("process", ""), lang) or "-"} / {meta.get("lot") or "-"}',
             f'- **{R("verdict", lang)}: {sit["label"]}** — {STATUS[lang][entry["status"]]}', '']
    if tx.get('summary'):
        lines += [tx['summary'], '']
    if det['defects']:
        lines += [f'| {R("c_no", lang)} | {R("c_type", lang)} | {R("c_level", lang)} | {R("c_conf", lang)} | '
                  f'{R("c_pos", lang)} | {R("c_reason", lang)} |', '|---|---|---|---|---|---|']
        for d in det['defects']:
            view = f'{d["view"]} ' if d.get('view') else ''
            lines.append(f'| {d["id"]} | {type_name(d["type"], kb_map.get(d["type"]), lang)} | '
                         f'{verdict_name(d["level"], lang)} | {d["conf"]:.2f} | {view}{pos_text(d["position"], lang)} | '
                         f'{reason_text(d, lang)} |')
        lines.append('')
    for key, label in (('cause_text', 'h_causes'), ('action_immediate', 'a_now'), ('action_prevent', 'a_prevent')):
        if tx.get(key):
            lines += [f'**{R(label, lang).lstrip("0123456789. ")}**: {tx[key]}', '']
    for d in tx.get('defects', []):
        if d.get('description'):
            lines.append(f'- {type_name(d["type"], kb_map.get(d["type"]), lang)}: {d["description"]}')
    imgs = [v['marked_image'] for v in det['views']] if det.get('views') else [det.get('marked_image')]
    lines += [''] + [f'![]({_rel(p, path.parent)})' for p in imgs if p]
    if result.get('report'):
        lines += ['', f'[{R("manager_title" if det["verdict"] == "reject" else "inspect_title", lang)}]'
                      f'({_rel(result["report"], path.parent)})']
    path.write_text('\n'.join(lines) + '\n', encoding='utf-8')
    return path


def load(out_root):
    """id 별 최신 기록 (처음 기록된 순서)."""
    path = Path(out_root) / JSONL
    latest = {}
    if path.exists():
        for line in path.read_text(encoding='utf-8').splitlines():
            if line.strip():
                e = json.loads(line)
                latest[e['id']] = e
    return list(latest.values())


def pending(out_root):
    return [e for e in load(out_root) if e['status'] == 'pending']


def upsert(result, out_root, status):
    """기록 추가/갱신 → md 작성 → 엑셀 다시 만들기. 엑셀이 열려 있어 저장 못 하면 경고 문자열을 돌려준다."""
    out_root = Path(out_root)
    e = make_entry(result, status, out_root)
    e['md'] = _rel(write_md(result, e, out_root), out_root)
    with open(out_root / JSONL, 'a', encoding='utf-8') as f:
        f.write(json.dumps(e, ensure_ascii=False) + '\n')
    return build_xlsx(out_root)


def _thumb(src, out_root):
    p = Path(src)
    t = Path(out_root) / 'work' / '_thumbs' / f'{p.stem}_t.png'
    if not t.exists():
        t.parent.mkdir(parents=True, exist_ok=True)
        im = Image.open(p)
        im.thumbnail((THUMB_H * 4, THUMB_H))
        im.save(t)
    return t


# ── 엑셀 대장 (영어) ─────────────────────────────────────────────────
# 기록에는 언어와 무관한 값(domain, verdict, rep, process, 결함별 class·level·reason_key)이 있고,
# 그 값으로 영어 문구를 만든다. 이 값이 없는 예전 기록은 한국어 값을 거꾸로 바꿔 쓴다.
_KO_VERDICT = {'정상': 'normal', '의심': 'review', '이상': 'reject'}
_KO_DOMAIN = {'PCB': 'pcb', '금속': 'msdd'}
_KO_PROCESS = {v: k for k, v in PROCESS_NAME['ko'].items() if k}
DOMAIN_EN = {'pcb': 'PCB', 'msdd': 'Metal'}
_OUT_ROOT = [Path('outputs')]        # build_xlsx 가 정한다 (예전 기록의 결과 파일 경로 기준)


def _kb(domain, cls):
    return TAXONOMIES.get(domain, {}).get('defects', {}).get(cls)


def _en_entry(e):
    dom = e.get('domain') or _KO_DOMAIN.get(e.get('도메인'), 'pcb')
    verdict = e.get('verdict') or _KO_VERDICT.get(e.get('판정'), 'normal')
    rep = e.get('rep')
    if not rep and e.get('대표 유형'):  # 예전 기록: 한국어 이름 → 클래스
        rep = next((c for c, k in TAXONOMIES[dom]['defects'].items() if k.get('name_ko') == e['대표 유형']), '')
    process = e.get('process') or _KO_PROCESS.get(e.get('공정'), '')
    old = {}                         # 예전 기록: 판정 근거 키를 저장된 검사 결과에서 다시 읽는다
    if any(not d.get('reason_key') for d in e['defects']) and e.get('result_json'):
        try:
            res = json.loads((_OUT_ROOT[0] / e['result_json']).read_text(encoding='utf-8'))
            old = {d['id']: d for d in res['detection']['defects']}
        except Exception:  # noqa: BLE001 — 결과 파일이 없으면 한국어 근거 그대로
            pass
    defects = []
    for d in e['defects']:
        if not d.get('reason_key') and d['no'] in old:
            d = {**d, 'reason_key': old[d['no']].get('reason_key'), 'reason_args': old[d['no']].get('reason_args', {})}
        level = d.get('level') or _KO_VERDICT.get(d.get('판정'), 'review')
        defects.append({
            'No': d['no'], 'View': d.get('사진', ''), 'Type': type_name(d['클래스'], _kb(dom, d['클래스']), 'en'),
            'Class': d['클래스'], 'Verdict': verdict_name(level, 'en'), 'v': level, 'Confidence': d['신뢰도'],
            'Position': pos_text(d['위치'], 'en'), 'Size (px)': d['크기(px)'], 'Area (%)': d['면적(%)'],
            'Basis': reason_text(d, 'en') if d.get('reason_key') else d.get('판정 근거', '')})
    return {
        'Time': e['일시'], 'Domain': DOMAIN_EN.get(dom, dom), 'Target': e['대상'], 'Photos': e['사진 수'],
        'Part no.': e.get('품번', ''), 'Process': process_name(process, 'en') if process else '', 'Lot': e.get('로트', ''),
        'Verdict': verdict_name(verdict, 'en'), 'v': verdict,
        'Main defect': type_name(rep, _kb(dom, rep), 'en') if rep else '',
        'Abnormal': e['이상'], 'Suspect': e['의심'], 'Worker edits': e.get('작업자 수정', 0),
        'Status': STATUS['en'].get(e['status'], e.get('처리 상태', '')), 'status': e['status'],
        'Report': e.get('보고서', ''), 'Record (md)': e.get('md', ''), 'image': e.get('image'), 'defects': defects}


def build_xlsx(out_root):
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    out_root = Path(out_root)
    _OUT_ROOT[0] = out_root
    rows = [_en_entry(e) for e in load(out_root)]
    wb = Workbook()

    def sheet(ws, headers, data, widths):
        ws.append(headers)
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = PatternFill('solid', fgColor='E0E0E0')
        for r in data:
            ws.append(r)
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions

    # 1. Inspections (검사 1건 = 1줄)
    ws = wb.active
    ws.title = 'Inspections'
    cols = ['Photo', 'Time', 'Domain', 'Target', 'Photos', 'Part no.', 'Process', 'Lot', 'Verdict', 'Main defect',
            'Abnormal', 'Suspect', 'Worker edits', 'Status', 'Report', 'Record (md)']
    sheet(ws, cols, [[''] + [e.get(c, '') for c in cols[1:]] for e in reversed(rows)],
          [12, 19, 7, 34, 7, 12, 14, 12, 10, 16, 9, 8, 12, 36, 40, 40])
    v_col = cols.index('Verdict') + 1
    for i, e in enumerate(reversed(rows), start=2):
        ws.row_dimensions[i].height = THUMB_H * 0.78
        ws.cell(i, v_col).fill = PatternFill('solid', fgColor=FILL[e['v']])
        for c in ws[i]:
            c.alignment = Alignment(vertical='center', wrap_text=False)
        try:
            ws.add_image(XLImage(str(_thumb(e['image'], out_root))), f'A{i}')
        except Exception:  # noqa: BLE001 — 사진이 지워졌으면 썸네일 없이
            pass
        for col_name in ('Report', 'Record (md)'):
            if e.get(col_name):
                cell = ws.cell(i, cols.index(col_name) + 1)
                cell.hyperlink = e[col_name]
                cell.font = Font(color='1565C0', underline='single')

    # 2. Defects (결함 1개 = 1줄)
    dcols = ['Time', 'Target', 'Part no.', 'Process', 'No', 'View', 'Type', 'Class', 'Verdict', 'Confidence',
             'Position', 'Size (px)', 'Area (%)', 'Basis']
    data, fills = [], []
    for e in reversed(rows):
        for d in e['defects']:
            data.append([e['Time'], e['Target'], e['Part no.'], e['Process']] + [d.get(c, '') for c in dcols[4:]])
            fills.append(d['v'])
    ws2 = wb.create_sheet('Defects')
    sheet(ws2, dcols, data, [19, 34, 12, 14, 5, 12, 16, 13, 10, 10, 14, 12, 9, 48])
    for i, v in enumerate(fills, start=2):
        ws2.cell(i, dcols.index('Verdict') + 1).fill = PatternFill('solid', fgColor=FILL[v])

    # 3. Summary (날짜별 · 결함 유형별 건수)
    ws3 = wb.create_sheet('Summary')
    verdicts = ['normal', 'review', 'reject']
    by_day = Counter((e['Time'][:10], e['v']) for e in rows)
    days = sorted({d for d, _ in by_day}, reverse=True)
    ws3.append(['Date', *(verdict_name(v, 'en') for v in verdicts), 'Total', 'Waiting for worker'])
    for d in days:
        n = [by_day[(d, v)] for v in verdicts]
        wait = sum(e['Time'][:10] == d and e['status'] == 'pending' for e in rows)
        ws3.append([d, *n, sum(n), wait])
    ws3.append([])
    by_type = Counter((x['Type'], x['v']) for e in rows for x in e['defects'])
    ws3.append(['Defect type', 'Suspect', 'Abnormal', 'Total'])
    for tp in sorted({tp for tp, _ in by_type}):
        ws3.append([tp, by_type[(tp, 'review')], by_type[(tp, 'reject')], by_type[(tp, 'review')] + by_type[(tp, 'reject')]])
    for row in ws3.iter_rows():
        if row and row[0].value in ('Date', 'Defect type'):
            for c in row:
                c.font = Font(bold=True)
                c.fill = PatternFill('solid', fgColor='E0E0E0')
    for col, w in zip('ABCDEF', (16, 10, 10, 10, 8, 20)):
        ws3.column_dimensions[col].width = w

    path = out_root / XLSX
    try:
        wb.save(path)
        return None
    except PermissionError:          # 엑셀에서 열어 둔 상태
        alt = out_root / f'{Path(XLSX).stem}_{time.strftime("%H%M%S")}.xlsx'
        wb.save(alt)
        return f'{XLSX} is open in Excel — saved as {alt.name}'
