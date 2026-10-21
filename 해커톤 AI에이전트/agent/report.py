"""보고서 생성 → Word(.docx). 판정에 따라 업계에서 흔히 쓰는 두 양식으로 만든다. 언어는 result['lang'] (ko / vi / en).

- 이상(reject): 부적합 보고서 NCR (Nonconformance Report, ISO 9001 부적합 관리)
    부적합 내용(5W2H) → 증거 사진 → 결함 상세 → 긴급 조치 → 처리 결정 → 원인 후보 → 시정·재발 방지 → 8D 연계 → 승인
    NCR 의 앞부분이 8D 의 D2(문제 기술)·D3(봉쇄 조치)·D4(원인 후보)에 해당한다. 반복·치명 결함이면 8D 발행을 권고한다.
- 의심(review): 재검사 의뢰서 (Hold) — 부품을 보류하고 재검사 결과로 Hold 해제(합격) 또는 NCR 발행(불합격)
정상은 보고서 없이 이력만 남긴다 (pipeline).
"""
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from agent.detect import thresholds
from agent.knowledge import localize
from agent.i18n import DOMAIN, SEVERITY, R, pos_text, process_name, reason_text, situation_text, type_name, verdict_name

FONT = '맑은 고딕'                   # 한글·베트남어 모두 표시됨
VERDICT_COLOR = {'normal': RGBColor(0x2E, 0x7D, 0x32), 'review': RGBColor(0xE6, 0x8A, 0x00),
                 'reject': RGBColor(0xC6, 0x28, 0x28)}
SEVERITY_KO = SEVERITY['ko']
SITUATION_KO = {'normal': '정상', 'review': '의심', 'reject': '이상'}
SEV_ORDER = {'critical': 0, 'major': 1, 'minor': 2}
DOC_PREFIX = {'reject': 'NCR', 'review': 'RIR'}      # 문서 번호 머리 (NCR / Re-inspection Request)
MAX_CROPS = 6
LABEL_FILL = 'EEF1F4'


def _set_font(doc):
    st = doc.styles['Normal']
    st.font.name = FONT
    st.font.size = Pt(10)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    for name in ('Title', 'Heading 1', 'Heading 2'):
        s = doc.styles[name]
        s.font.name = FONT
        s.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)


def _shade(cell, fill=LABEL_FILL):
    pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    pr.append(shd)


def _table(doc, rows, header=True, widths=None, size=None, label_cols=()):
    """표. header=True 면 첫 행 굵게·음영, label_cols 의 열은 항목 이름 칸(굵게·음영)."""
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = 'Table Grid'
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            cell = t.cell(i, j)
            cell.text = '' if v is None else str(v)
            is_label = (header and i == 0) or j in label_cols
            for r in cell.paragraphs[0].runs:
                r.bold = is_label
                if size:
                    r.font.size = Pt(size)
            if is_label:
                _shade(cell)
            if widths:
                cell.width = widths[j]
    return t


class _Ctx:
    """보고서 한 건을 만드는 동안 쓰는 값 모음."""

    def __init__(self, result):
        self.result = result
        self.det, self.kb, self.text = result['detection'], result['knowledge'], result['text']
        self.meta = result.get('meta') or {}
        self.lang = result.get('lang', 'ko')
        self.kb_map = {k['type']: k for k in self.kb}
        raw = result['situation']
        self.sit = situation_text(raw, self.det['verdict'], self.lang)
        self.extra = {k: situation_text(raw[k], k, self.lang) for k in ('part', 'surface_wide', 'multi_type') if k in raw}
        self.rep = raw.get('representative')
        self.th = thresholds(self.det['domain'])

    def L(self, key, **kw):
        return R(key, self.lang, **kw)

    def name(self, t):
        return type_name(t, self.kb_map.get(t), self.lang)

    def sev_key(self, t):
        return self.kb_map.get(t, {}).get('severity')

    def sev(self, t):
        return SEVERITY[self.lang].get(self.sev_key(t), '-')

    def vname(self, v):
        return verdict_name(v, self.lang)


def doc_number(result):
    """문서 번호 (NCR-20261017-143015 / RIR-…). 미리보기와 저장본이 같은 번호가 되도록 결과에 한 번만 붙인다."""
    v = result['detection']['verdict']
    nos = result.setdefault('doc_no', {})
    if v not in nos:
        nos[v] = f'{DOC_PREFIX[v]}-{datetime.now():%Y%m%d-%H%M%S}'
    return nos[v]


def _new_doc(title, note):
    doc = Document()
    _set_font(doc)
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(1.8)
        s.top_margin = s.bottom_margin = Cm(1.6)
    doc.add_heading(title, level=0)
    r = doc.add_paragraph().add_run(note)
    r.italic = True
    r.font.size = Pt(9)
    return doc


def _banner(doc, text, color):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = color
    return p


def _source(c):
    det = c.det
    if det.get('views'):
        return f'{det.get("part") or "-"} · ' + ', '.join(f'{v["view"]}: {c.vname(v["verdict"])}' for v in det['views'])
    return Path(det['image']).name


def _header(doc, c, no_label, no, status_label, status_value):
    """문서 머리 표: 번호·일자 / 품번·공정 / 로트·대상 / 사진·상태 (2쌍씩 4열)."""
    L, m = c.L, c.meta
    rows = [
        [L(no_label), no, L('date'), datetime.now().strftime('%Y-%m-%d %H:%M')],
        [L('m_part_no'), m.get('part_no') or '-', L('m_process'), process_name(m.get('process'), c.lang) if m.get('process') else '-'],
        [L('m_lot'), m.get('lot') or '-', L('target'), DOMAIN[c.lang].get(c.det['domain'], c.det['domain'])],
        [L('photos' if c.det.get('views') else 'file'), _source(c), L(status_label), status_value],
    ]
    _table(doc, rows, header=False, widths=[Cm(2.9), Cm(5.8), Cm(2.9), Cm(5.8)], label_cols=(0, 2))


def _photos(doc, c, heading):
    det = c.det
    doc.add_heading(heading, level=1)
    if det.get('views'):             # 부품 단위: 방향별 표시 이미지 3열
        vs = det['views']
        t = doc.add_table(rows=2 * ((len(vs) + 2) // 3), cols=3)
        for i, v in enumerate(vs):
            r0, col = 2 * (i // 3), i % 3
            t.cell(r0, col).text = f'{v["view"]} — {c.vname(v["verdict"])} ({c.L("n_items", n=v["n_defects"])})'
            t.cell(r0 + 1, col).paragraphs[0].add_run().add_picture(v['marked_image'] or v['image'], width=Cm(5.3))
    else:
        t = doc.add_table(rows=2, cols=2)
        t.cell(0, 0).text, t.cell(0, 1).text = c.L('original'), c.L('marked')
        t.cell(1, 0).paragraphs[0].add_run().add_picture(det['image'], width=Cm(8))
        if det.get('marked_image'):
            t.cell(1, 1).paragraphs[0].add_run().add_picture(det['marked_image'], width=Cm(8))
        for cell in t.rows[0].cells:
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER


def _size(d):
    s = f'{d["size_px"][0]}×{d["size_px"][1]}'
    return s + f' ({d["area_ratio"] * 100:.1f}%)' if 'area_ratio' in d else s


def _crops(doc, c, caption):
    crops = [d for d in c.det['defects'] if d.get('crop')][:MAX_CROPS]
    if not crops:
        return
    doc.add_paragraph(caption)
    ct = doc.add_table(rows=2, cols=len(crops))
    for j, d in enumerate(crops):
        ct.cell(0, j).paragraphs[0].add_run().add_picture(d['crop'], height=Cm(2.5))
        ct.cell(1, j).text = f'#{d["id"]} {c.name(d["type"])} {d["conf"]:.2f} {c.vname(d["level"])}'


def _references(doc, c, L):
    """본문 문장에서 [S#] 로 인용된 웹 수집 문서만 출처와 함께 나열한다."""
    import re
    cited = set(re.findall(r'\[(S\d+)\]', ' '.join(str(c.text.get(k, '')) for k in ('summary', 'cause_text', 'action_immediate', 'action_prevent'))))
    srcs = [s for s in c.text.get('sources', []) if s['id'] in cited]
    if srcs:
        p = doc.add_paragraph()
        p.add_run(L('refs') + ': ').bold = True
        p.add_run('; '.join(f'[{s["id"]}] {s["title"]} ({s["url"]}, {s["license"]})' for s in srcs)).font.size = Pt(8.5)


def _situation_notes(doc, c):
    for key, s in c.extra.items():
        types = f' ({c.L("applies_to")}: {", ".join(c.name(x) for x in s["types"])})' if 'types' in s else ''
        doc.add_paragraph(f'{s["label"]}{types}: {s["action"]}')
    if c.lang == 'ko':               # 지식베이스 원문 메모는 한국어만
        for k in c.kb:
            if k.get('surface_note') and k['type'] in c.extra.get('surface_wide', {}).get('types', []):
                doc.add_paragraph(f'{k["name_ko"]}: {k["surface_note"]}')


def _defect_rows(c, cols):
    has_view = 'view' in c.det['defects'][0]
    return has_view, [[c.L('c_no'), *([c.L('c_view')] if has_view else []), *cols]]


def _problem(doc, c):
    """부적합 내용 — 5W2H (8D D2 문제 기술)."""
    L, det = c.L, c.det
    defects = det['defects']
    n_rej = sum(d['level'] == 'reject' for d in defects)
    rej = [d for d in defects if d['level'] == 'reject'] or defects
    types = list(dict.fromkeys(d['type'] for d in rej))
    where = ', '.join(dict.fromkeys(pos_text(d['position'], c.lang) for d in rej))
    if det.get('views'):
        where = f'{det.get("part") or "-"} — ' + ', '.join(dict.fromkeys(d.get('view', '') for d in rej)) + f' / {where}'
    how = L('w_how_v', conf=max(d['conf'] for d in defects))
    if det.get('tiles', 1) > 1:
        how += ' · ' + L('tiled', n=det['tiles'])
    rows = [
        [L('w_what'), ', '.join(f'{c.name(t)} ({c.sev(t)})' for t in types)],
        [L('w_where'), where],
        [L('w_when'), datetime.now().strftime('%Y-%m-%d %H:%M')
         + (f' · {process_name(c.meta["process"], c.lang)}' if c.meta.get('process') else '')],
        [L('w_who'), L('w_who_v')],
        [L('w_how'), how],
        [L('w_howmany'), L('count_val', n=len(defects), r=n_rej, s=len(defects) - n_rej)],
        [L('w_req'), L('w_req_v', review=c.th['conf_review'], reject=c.th['conf_reject'])],
    ]
    _table(doc, rows, header=False, widths=[Cm(4.2), Cm(13.2)], label_cols=(0,))
    if c.text.get('summary'):
        doc.add_paragraph(c.text['summary'])


def _needs_8d(c):
    """8D 발행 권고 사유 (치명 결함 / 여러 유형 동시 / 전면 분포)."""
    why = []
    if any(c.sev_key(d['type']) == 'critical' for d in c.det['defects'] if d['level'] == 'reject'):
        why.append(c.L('d8_why_critical'))
    if 'multi_type' in c.extra:
        why.append(c.L('d8_why_multi'))
    if 'surface_wide' in c.extra:
        why.append(c.L('d8_why_surface'))
    return why


def _ncr(result):
    c = _Ctx(result)
    L = c.L
    doc = _new_doc(L('manager_title'), L('manager_note'))
    _header(doc, c, 'ncr_no', doc_number(result), 'ncr_status', L('ncr_status_v'))
    worst = min((c.sev_key(d['type']) for d in c.det['defects'] if d['level'] == 'reject'),
                key=lambda s: SEV_ORDER.get(s, 9), default=None)
    _banner(doc, f'{L("verdict")}: {c.sit["label"]}'
            + (f'  ·  {L("c_sev")}: {SEVERITY[c.lang].get(worst)}' if worst else ''), VERDICT_COLOR['reject'])

    doc.add_heading(L('n_h_problem'), level=1)
    _problem(doc, c)

    _photos(doc, c, L('n_h_photo'))
    _crops(doc, c, L('crops'))

    doc.add_heading(L('n_h_defects'), level=1)
    has_view, rows = _defect_rows(c, [L('c_type'), L('c_sev'), L('c_pos'), L('c_size'), L('c_conf'), L('c_level'), L('c_reason')])
    for d in c.det['defects']:
        rows.append([d['id'], *([d['view']] if has_view else []), c.name(d['type']), c.sev(d['type']), pos_text(d['position'], c.lang),
                     _size(d), f'{d["conf"]:.2f}', c.vname(d['level']), reason_text(d, c.lang)])
    _table(doc, rows, size=8)
    for d in c.text.get('defects', []):
        doc.add_paragraph(f'{c.name(d["type"])}: {d["description"]}', style='List Bullet')

    doc.add_heading(L('n_h_contain'), level=1)
    if c.text.get('action_immediate'):
        doc.add_paragraph(c.text['action_immediate'])
    _table(doc, [[L('t_action'), L('t_owner'), L('t_date')],
                 [L('contain_1'), '', ''], [L('contain_2'), '', ''], [L('contain_3'), '', '']],
           widths=[Cm(11.4), Cm(3.2), Cm(2.8)])

    doc.add_heading(L('n_h_dispo'), level=1)
    _table(doc, [[L('dispo'), L('dispo_v')], [L('dispo_qty'), ''], [L('dispo_by'), '']],
           header=False, widths=[Cm(4.2), Cm(13.2)], label_cols=(0,))

    doc.add_heading(L('n_h_cause'), level=1)
    if c.text.get('cause_text'):
        doc.add_paragraph(c.text['cause_text'])
    # 지식베이스 원문 (한국어) / 번역본 (translations_<언어>.yaml — 번역 없으면 건너뜀)
    for k in c.kb:
        if c.lang == 'ko' or k.get(c.lang):
            doc.add_paragraph(f'{c.name(k["type"])} ({c.sev(k["type"])}): ' + ' / '.join(localize(k, c.lang)['causes']), style='List Bullet')
    _references(doc, c, L)
    note = doc.add_paragraph().add_run(L('cause_note'))
    note.italic = True
    note.font.size = Pt(9)

    doc.add_heading(L('n_h_corrective'), level=1)
    rows = [[L('a_kind'), L('a_content')], [L('a_now'), c.text.get('action_immediate', '')],
            [L('a_prevent'), c.text.get('action_prevent', '')]]
    for k in c.kb:
        if c.lang != 'ko' and not k.get(c.lang):
            continue
        a = localize(k, c.lang).get('actions', {})
        rows.append([f'{L("a_ref")} {c.name(k["type"])}',
                     f'{L("a_ref_now")}: {a.get("immediate", "-")}\n{L("a_ref_check")}: {a.get("check", "-")}\n'
                     f'{L("a_ref_prevent")}: {a.get("prevent", "-")}'])
    _table(doc, rows, widths=[Cm(3.5), Cm(13.9)])
    _situation_notes(doc, c)

    doc.add_heading(L('n_h_8d'), level=1)
    why = _needs_8d(c)
    _banner(doc, L('d8_yes', why=', '.join(why)) if why else L('d8_no'),
            VERDICT_COLOR['reject'] if why else RGBColor(0x55, 0x5F, 0x6A)).runs[0].font.size = Pt(10.5)
    _table(doc, [[L('d8_step'), L('d8_state')],
                 ['D1 ' + L('d8_d1'), L('d8_todo')],
                 ['D2 ' + L('d8_d2'), L('d8_done', sec=1)],
                 ['D3 ' + L('d8_d3'), L('d8_partial', sec=4)],
                 ['D4 ' + L('d8_d4'), L('d8_input', sec=6)],
                 ['D5 ' + L('d8_d5'), L('d8_input', sec=7)],
                 ['D6 ' + L('d8_d6'), L('d8_todo')],
                 ['D7 ' + L('d8_d7'), L('d8_input', sec=7)],
                 ['D8 ' + L('d8_d8'), L('d8_todo')]],
           widths=[Cm(6.2), Cm(11.2)], size=9)
    _table(doc, [[L('d8_issue'), L('d8_issue_v')]], header=False, widths=[Cm(4.2), Cm(13.2)], label_cols=(0,))

    doc.add_heading(L('n_h_sign'), level=1)
    _table(doc, [[L('w_confirm'), L('w_confirm_v')], [L('w_edit'), '']],
           header=False, widths=[Cm(4.2), Cm(13.2)], label_cols=(0,))
    doc.add_paragraph()
    _table(doc, [[L('s_role'), L('s_name'), L('s_sign'), L('s_date')],
                 [L('s_author'), '', '', ''], [L('s_review'), '', '', ''], [L('s_approve'), '', '', ''],
                 [L('s_close'), '', '', '']],
           widths=[Cm(5.4), Cm(4.4), Cm(4.0), Cm(3.6)])
    return doc


def _hold(result):
    c = _Ctx(result)
    L = c.L
    rules = c.th.get('rules', {})
    doc = _new_doc(L('inspect_title'), L('inspect_note'))
    _banner(doc, L('hold_banner'), VERDICT_COLOR['review'])
    _header(doc, c, 'rir_no', doc_number(result), 'hold_status', L('hold_status_v'))

    doc.add_heading(L('r_h_reason'), level=1)
    run = doc.add_paragraph().add_run(f'{L("verdict")}: {c.sit["label"]}'
                                      + (f'  →  {c.sit["output"]}' if c.sit.get('output') else ''))
    run.bold = True
    run.font.color.rgb = VERDICT_COLOR['review']
    doc.add_paragraph(L('hold_rule', review=c.th['conf_review'], reject=c.th['conf_reject']))
    if c.text.get('summary'):
        doc.add_paragraph(c.text['summary'])

    _photos(doc, c, L('r_h_photo'))
    _crops(doc, c, L('crops2'))

    doc.add_heading(L('r_h_items'), level=1)
    has_view, rows = _defect_rows(c, [L('c_stype'), L('c_pos'), L('c_size'), L('c_conf'), L('c_sreason'), L('c_how'), L('c_check')])
    for d in c.det['defects']:
        rule = rules.get(d['type'], {})
        if c.lang == 'ko':
            how = rule.get('note') or c.kb_map.get(d['type'], {}).get('actions', {}).get('immediate') or L('how_default')
        else:
            how = (rule.get(f'note_{c.lang}')
                   or (localize(c.kb_map[d['type']], c.lang).get('actions', {}).get('immediate') if c.kb_map.get(d['type'], {}).get(c.lang) else None)
                   or L('how_default'))
        rows.append([d['id'], *([d['view']] if has_view else []), c.name(d['type']), pos_text(d['position'], c.lang), _size(d),
                     f'{d["conf"]:.2f}', reason_text(d, c.lang), how, L('check_box')])
    _table(doc, rows, size=8)

    if c.text.get('defects'):
        doc.add_heading(L('r_h_ai'), level=1)
        for d in c.text['defects']:
            doc.add_paragraph(f'{c.name(d["type"])}: {d["description"]}', style='List Bullet')

    doc.add_heading(L('r_h_result'), level=1)
    _situation_notes(doc, c)
    _table(doc, [[L('r_method'), L('r_method_v')], [L('f_final'), L('f_final_v')], [L('f_memo'), ''],
                 [L('f_sign'), ''], [L('r_release'), '']],
           header=False, widths=[Cm(4.2), Cm(13.2)], label_cols=(0,))
    if c.sit.get('note'):
        doc.add_paragraph(c.sit['note']).runs[0].font.size = Pt(9)
    return doc


def build_report(result, out_path, part_info=''):
    """이상 → NCR, 의심 → 재검사 의뢰서(Hold). part_info 는 예전 호출과의 호환용 (품번·공정·로트는 result['meta'])."""
    verdict = result['detection']['verdict']
    if verdict == 'normal':
        return None
    doc = (_ncr if verdict == 'reject' else _hold)(result)
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return str(out_path)
